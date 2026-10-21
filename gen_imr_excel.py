# Gera a versao Excel (.xlsx) do IMR com as mesmas funcionalidades do formulario HTML.
# Uso: python gen_imr_excel.py "Angels" "04/24" "CBAE" "IMR Angels - 04-24 - CBAE.xlsx"
#
# Replica do formulario "Angels - 04-24 - CBAE.html": 60 periodos (abas), avaliacao de
# desempenho com pesos e notas 0/1/3, nota final, conceito e faixa de ajuste de pagamento,
# postos de trabalho com desconto por hora sem cobertura, instrumento de cobranca,
# resumo financeiro (30 avos), reembolso de creche com saldo acumulado entre os meses,
# assinaturas, justificativas e impressao em 1 folha A4.
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

NUM_PERIODS = 60
NP = 15          # linhas de posto disponiveis por aba
HORAS_MES = 220  # horas contratadas por posto/mes (igual ao formulario web)

# ── CORES (iguais ao formulario HTML) ────────────────────────────
AZUL_ESC, AZUL = '1F3864', '2E74B5'
AZUL_CLR, AZUL_TOT = 'D6E4F7', 'DCE6F1'
VERDE, AMARELO, VERMELHO, LARANJA = 'E8F5E9', 'FFFDE7', 'FFEBEE', 'FFF8E1'
VERDE_ESC, VERDE_IC, VERDE_ICH = '556B2F', 'F9FCE8', '6B8E23'

FONTE = 'Arial'
F_TXT  = Font(name=FONTE, size=10)
F_B    = Font(name=FONTE, size=10, bold=True)
F_IN   = Font(name=FONTE, size=10, color='0000FF')               # entrada manual
F_LINK = Font(name=FONTE, size=10, color='008000')               # vem da aba Dados
F_TIT  = Font(name=FONTE, size=11, bold=True, color='FFFFFF')
F_HDR  = Font(name=FONTE, size=9, bold=True, color='FFFFFF')
F_GRP  = Font(name=FONTE, size=10, bold=True, color=AZUL_ESC)
F_RED  = Font(name=FONTE, size=9, italic=True, bold=True, color='C00000')
F_SM   = Font(name=FONTE, size=8, italic=True, color='777777')

FILL_TIT = PatternFill('solid', fgColor=AZUL_ESC)
FILL_HDR = PatternFill('solid', fgColor=AZUL)
FILL_GRP = PatternFill('solid', fgColor=AZUL_CLR)
FILL_TOT = PatternFill('solid', fgColor=AZUL_TOT)
FILL_CAL = PatternFill('solid', fgColor=VERDE)
FILL_IN  = PatternFill('solid', fgColor=AMARELO)
FILL_RED = PatternFill('solid', fgColor=VERMELHO)
FILL_ADM = PatternFill('solid', fgColor=LARANJA)
FILL_IC  = PatternFill('solid', fgColor=VERDE_IC)
FILL_ICH = PatternFill('solid', fgColor=VERDE_ICH)
FILL_APR = PatternFill('solid', fgColor='C8E6C9')

_thin = Side(style='thin', color='999999')
BOX = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

MOEDA, NUM2, PCT, INT = 'R$ #,##0.00', '#,##0.00', '0%', '#,##0'
DATA_BR, MES_ANO = 'dd/mm/yyyy', '[$-416]mmmm"/"yyyy'

# ── CRITERIOS: (id, nome, peso, grupo, e_do_fiscal_administrativo) ──
CRITERIOS = [
    ('g1_1', 'Cumprimento das Atividades', 0.40, 1, False),
    ('g1_2', 'Qualificação / Atendimento ao Público / Postura', 0.35, 1, False),
    ('g1_3', 'Uniformes e Identificação', 0.25, 1, False),
    ('g2_1', 'Especificação Técnica dos Serviços', 0.40, 2, False),
    ('g2_2', 'Equipamentos e Acessórios', 0.20, 2, False),
    ('g2_3', 'Atendimento às Ocorrências', 0.40, 2, False),
    ('g3_1', 'Periodicidade da Supervisão', 0.20, 3, False),
    ('g3_2', 'Gerenciamento das Atividades Operacionais', 0.30, 3, False),
    ('g3_3', 'Atendimento às Solicitações', 0.25, 3, False),
    ('g3_4', 'Salários, Benefícios e Obrigações Trabalhistas (¹ Fiscal Adm.)', 0.25, 3, True),
]
GRUPO_NOME = {
    1: '▶ GRUPO 1 – DESEMPENHO PROFISSIONAL',
    2: '▶ GRUPO 2 – DESEMPENHO DAS ATIVIDADES',
    3: '▶ GRUPO 3 – GERENCIAMENTO',
}
# Faixa de ajuste em ordem CRESCENTE de nota minima (exigido pelo MATCH aproximado)
FAIXA = [
    (0.00, 'F', 0.75, 'Abaixo de 5,40'),
    (5.40, 'E', 0.80, 'Entre 5,84 e 5,40'),
    (5.85, 'D', 0.85, 'Entre 6,74 e 5,85'),
    (6.75, 'C', 0.90, 'Entre 7,64 e 6,75'),
    (7.65, 'B', 0.95, 'Entre 8,09 e 7,65'),
    (8.10, 'A', 1.00, 'Entre 9,00 e 8,10'),
]

LARGURAS = {'A': 42, 'B': 26, 'C': 12, 'D': 15, 'E': 13, 'F': 13, 'G': 15,
            'H': 17, 'I': 17, 'J': 11}


# ── utilitarios de escrita ───────────────────────────────────────
def put(ws, ref, value, font=F_TXT, fill=None, fmt=None, align=None,
        border=True, wrap=False, unlocked=False, valign='center'):
    c = ws[ref]
    c.value = value
    c.font = font
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(horizontal=align, vertical=valign, wrap_text=wrap)
    if border:
        c.border = BOX
    if unlocked:
        c.protection = Protection(locked=False)
    return c


def merge(ws, a, b, fill=None):
    ws.merge_cells(a + ':' + b)
    for row in ws[a + ':' + b]:
        for c in row:
            c.border = BOX
            if fill:
                c.fill = fill


def banner(ws, row, texto, last='J', fill=FILL_TIT, font=F_TIT, height=17):
    merge(ws, 'A%d' % row, '%s%d' % (last, row), fill=fill)
    put(ws, 'A%d' % row, texto, font=font, fill=fill, align='left')
    ws.row_dimensions[row].height = height


def rotulo(ws, r, texto, col_ini='A', col_fim='E', wrap=True):
    merge(ws, '%s%d' % (col_ini, r), '%s%d' % (col_fim, r), fill=FILL_TOT)
    put(ws, '%s%d' % (col_ini, r), texto, font=F_B, fill=FILL_TOT, wrap=wrap)


# ═════════════════════════════════════════════════════════════════
# ABA DADOS – valores compartilhados por todas as abas de periodo
# ═════════════════════════════════════════════════════════════════
def aba_dados(wb, contratada, contrato, unidade):
    ws = wb.create_sheet('Dados')
    for col, w in [('A', 36), ('B', 32), ('C', 14), ('D', 20), ('E', 34), ('F', 10)]:
        ws.column_dimensions[col].width = w
    ws.sheet_view.showGridLines = False

    banner(ws, 1, 'DADOS DO CONTRATO – compartilhados pelas 60 abas de período', last='F')
    for i, (lbl, v) in enumerate([
            ('CONTRATO Nº:', contrato),
            ('UNIDADE:', unidade),
            ('CONTRATADA:', contratada),
            ('FISCAL OPERACIONAL:', ''),
            ('OBSERVAÇÃO DO CABEÇALHO:',
             'Atualizado conforme Valor Unitário Contratado inicial')]):
        r = 2 + i
        put(ws, 'A%d' % r, lbl, font=F_B, fill=FILL_TOT)
        merge(ws, 'B%d' % r, 'F%d' % r, fill=FILL_IN)
        put(ws, 'B%d' % r, v, font=F_IN, fill=FILL_IN, unlocked=True)

    put(ws, 'A7', 'PERÍODO INICIAL (1ª aba – P01):', font=F_B, fill=FILL_TOT)
    merge(ws, 'B7', 'F7', fill=FILL_IN)
    put(ws, 'B7', None, font=F_IN, fill=FILL_IN, fmt=MES_ANO, unlocked=True)
    merge(ws, 'A8', 'F8')
    put(ws, 'A8', 'Digite qualquer dia do mês inicial (ex.: 01/05/2024). As abas P01…P60 '
                  'recebem os meses seguintes automaticamente.', font=F_SM, wrap=True)

    # ── Postos contratados (baseline replicado em todas as abas) ──
    banner(ws, 10, 'POSTOS DE TRABALHO CONTRATADOS (replicados em todas as abas)', last='F')
    for i, h in enumerate(['Item', 'POSTO', 'Qtde Contratada',
                           'Valor Unit. Contratado (R$)', 'Observação / fonte do valor', '']):
        put(ws, '%s11' % get_column_letter(1 + i), h, font=F_HDR, fill=FILL_HDR,
            align='center', wrap=True)
    ws.row_dimensions[11].height = 28
    posto_row0 = 12
    for i in range(NP):
        r = posto_row0 + i
        exemplo = (i == 0)
        put(ws, 'A%d' % r, 1 if exemplo else None, font=F_IN, fill=FILL_IN,
            align='center', unlocked=True)
        put(ws, 'B%d' % r, 'Vigilante Diurno 12x36 (exemplo)' if exemplo else None,
            font=F_IN, fill=FILL_IN, unlocked=True)
        put(ws, 'C%d' % r, 2 if exemplo else None, font=F_IN, fill=FILL_IN,
            align='center', unlocked=True)
        put(ws, 'D%d' % r, 5996.89 if exemplo else None, font=F_IN, fill=FILL_IN,
            fmt=MOEDA, unlocked=True)
        put(ws, 'E%d' % r, 'LINHA DE EXEMPLO – substitua pelos postos e valores do '
                           'contrato' if exemplo else None, font=F_SM, wrap=True,
            unlocked=True)
        put(ws, 'F%d' % r, None)
    rt = posto_row0 + NP
    merge(ws, 'A%d' % rt, 'C%d' % rt, fill=FILL_TOT)
    put(ws, 'A%d' % rt, 'CONTRATADO MENSAL – TOTAL (conferência)', font=F_B, fill=FILL_TOT)
    put(ws, 'D%d' % rt, '=SUMPRODUCT(C%d:C%d,D%d:D%d)'
        % (posto_row0, rt - 1, posto_row0, rt - 1), font=F_B, fill=FILL_TOT, fmt=MOEDA)
    put(ws, 'E%d' % rt, None, fill=FILL_TOT)
    put(ws, 'F%d' % rt, None, fill=FILL_TOT)

    # ── Creche – valores padrao das abas ─────────────────────────
    rc = rt + 2
    banner(ws, rc, 'REEMBOLSO DE CRECHE – valores padrão das abas', last='F')
    creche_retido = 'D%d' % (rc + 1)
    creche_unit = 'D%d' % (rc + 2)
    for i, (lbl, v) in enumerate([('Valor Reembolso Creche Retido (R$) por mês', 0),
                                  ('Reembolso por dependente (R$)', 0)]):
        r = rc + 1 + i
        merge(ws, 'A%d' % r, 'C%d' % r, fill=FILL_TOT)
        put(ws, 'A%d' % r, lbl, font=F_B, fill=FILL_TOT)
        put(ws, 'D%d' % r, v, font=F_IN, fill=FILL_IN, fmt=MOEDA, unlocked=True)
        put(ws, 'E%d' % r, None)
        put(ws, 'F%d' % r, None)
    merge(ws, 'A%d' % (rc + 3), 'F%d' % (rc + 3))
    put(ws, 'A%d' % (rc + 3), 'Entram em todas as abas e podem ser ajustados mês a mês '
                              'dentro de cada aba.', font=F_SM, wrap=True)

    # ── Criterios e pesos ───────────────────────────────────────
    rp = rc + 5
    banner(ws, rp, 'CRITÉRIOS E PESOS DA AVALIAÇÃO', last='F')
    for i, h in enumerate(['ID', 'Critério de Avaliação', 'Peso (a)', 'Grupo', '', '']):
        put(ws, '%s%d' % (get_column_letter(1 + i), rp + 1), h, font=F_HDR, fill=FILL_HDR,
            align='center')
    peso_row0 = rp + 2
    for i, (cid, nome, peso, grp, _adm) in enumerate(CRITERIOS):
        r = peso_row0 + i
        put(ws, 'A%d' % r, cid, font=F_TXT, align='center')
        put(ws, 'B%d' % r, nome, font=F_IN, fill=FILL_IN, wrap=True, unlocked=True)
        put(ws, 'C%d' % r, peso, font=F_IN, fill=FILL_IN, fmt=NUM2, align='center',
            unlocked=True)
        put(ws, 'D%d' % r, grp, font=F_TXT, align='center')
        put(ws, 'E%d' % r, None)
        put(ws, 'F%d' % r, None)
    rsp = peso_row0 + len(CRITERIOS)
    merge(ws, 'A%d' % rsp, 'B%d' % rsp, fill=FILL_TOT)
    put(ws, 'A%d' % rsp, 'SOMA DOS PESOS (deve ser 3,00 = 1,00 por grupo)', font=F_B,
        fill=FILL_TOT)
    put(ws, 'C%d' % rsp, '=SUM(C%d:C%d)' % (peso_row0, rsp - 1), font=F_B, fill=FILL_TOT,
        fmt=NUM2, align='center')
    for L in 'DEF':
        put(ws, '%s%d' % (L, rsp), None, fill=FILL_TOT)

    # ── Faixa de ajuste de pagamento ────────────────────────────
    rf = rsp + 2
    banner(ws, rf, 'FAIXA DE AJUSTE DE PAGAMENTO', last='F')
    for i, h in enumerate(['Nota Mínima', 'Conceito', '% do Faturamento a ser Pago',
                           'Nota Média Mensal Apurada', '', '']):
        put(ws, '%s%d' % (get_column_letter(1 + i), rf + 1), h, font=F_HDR, fill=FILL_HDR,
            align='center', wrap=True)
    ws.row_dimensions[rf + 1].height = 26
    faixa_row0 = rf + 2
    for i, (nmin, conc, pct, txt) in enumerate(FAIXA):
        r = faixa_row0 + i
        put(ws, 'A%d' % r, nmin, font=F_IN, fill=FILL_IN, fmt=NUM2, align='center',
            unlocked=True)
        put(ws, 'B%d' % r, conc, font=F_B, align='center')
        put(ws, 'C%d' % r, pct, font=F_IN, fill=FILL_IN, fmt=PCT, align='center',
            unlocked=True)
        put(ws, 'D%d' % r, txt, font=F_TXT, align='center')
        put(ws, 'E%d' % r, None)
        put(ws, 'F%d' % r, None)
    merge(ws, 'A%d' % (faixa_row0 + len(FAIXA)), 'F%d' % (faixa_row0 + len(FAIXA)))
    put(ws, 'A%d' % (faixa_row0 + len(FAIXA)),
        'Ordem crescente de nota mínima: é assim que as abas localizam o conceito e o % '
        'a liberar (INDEX/MATCH). Não reordene as linhas.', font=F_SM, wrap=True)

    ws.freeze_panes = 'A2'
    return {'posto_row0': posto_row0, 'peso_row0': peso_row0, 'faixa_row0': faixa_row0,
            'faixa_row1': faixa_row0 + len(FAIXA) - 1,
            'creche_retido': creche_retido, 'creche_unit': creche_unit}


# ═════════════════════════════════════════════════════════════════
# ABA DE PERIODO – uma por mês (P01…P60)
# ═════════════════════════════════════════════════════════════════
def aba_periodo(wb, n, D, nome_aba, aba_anterior):
    ws = wb.create_sheet(nome_aba)
    for col, w in LARGURAS.items():
        ws.column_dimensions[col].width = w
    ws.sheet_view.showGridLines = False

    fx0, fx1, pz0 = D['faixa_row0'], D['faixa_row1'], D['peso_row0']

    # ── CABEÇALHO ────────────────────────────────────────────────
    banner(ws, 1, 'UFRJ – PRÓ-REITORIA DE GESTÃO E GOVERNANÇA – PR6   |   IMR')
    merge(ws, 'A2', 'J2')
    put(ws, 'A2', 'INSTRUMENTO DE MEDIÇÃO DE RESULTADO',
        font=Font(name=FONTE, size=12, bold=True, color=AZUL_ESC), align='center')
    ws.row_dimensions[2].height = 19
    merge(ws, 'A3', 'J3')
    put(ws, 'A3', '=IF(Dados!$B$6="","",Dados!$B$6)', font=F_RED, align='center')

    # ── IDENTIFICAÇÃO ───────────────────────────────────────────
    def ident(linha, lbl_esq, ref_esq, lbl_dir, ref_dir):
        put(ws, 'A%d' % linha, lbl_esq, font=F_B, fill=FILL_TOT)
        merge(ws, 'B%d' % linha, 'D%d' % linha)
        put(ws, 'B%d' % linha, '=IF(Dados!$B${0}="","",Dados!$B${0})'.format(ref_esq),
            font=F_LINK)
        merge(ws, 'E%d' % linha, 'F%d' % linha, fill=FILL_TOT)
        put(ws, 'E%d' % linha, lbl_dir, font=F_B, fill=FILL_TOT)
        merge(ws, 'G%d' % linha, 'J%d' % linha)
        put(ws, 'G%d' % linha, '=IF(Dados!$B${0}="","",Dados!$B${0})'.format(ref_dir),
            font=F_LINK)

    ident(5, 'CONTRATO Nº:', 2, 'UNIDADE:', 3)
    ident(6, 'CONTRATADA:', 4, 'FISCAL OPERACIONAL:', 5)

    put(ws, 'A7', 'AVALIAÇÃO REFERENTE AO PERÍODO:', font=F_B, fill=FILL_TOT)
    merge(ws, 'B7', 'D7')
    put(ws, 'B7', '=IF(Dados!$B$7="","(informe o período inicial na aba Dados)",'
                  'EDATE(Dados!$B$7,%d))' % (n - 1),
        font=Font(name=FONTE, size=11, bold=True, color='008000'), fmt=MES_ANO,
        align='center')
    merge(ws, 'E7', 'F7', fill=FILL_TOT)
    put(ws, 'E7', 'PERÍODO (ABA) Nº:', font=F_B, fill=FILL_TOT)
    merge(ws, 'G7', 'J7')
    put(ws, 'G7', '%d de %d' % (n, NUM_PERIODS), font=F_B, align='center')

    merge(ws, 'A8', 'J8')
    put(ws, 'A8', '⚠ Para notas 0 e 1 o campo JUSTIFICATIVAS (fim da aba) é de preenchimento '
                  'obrigatório.   |   Contrato, Unidade, Contratada e Postos são '
                  'compartilhados por todas as abas (aba Dados).', font=F_RED, wrap=True)
    ws.row_dimensions[8].height = 24

    # ── AVALIAÇÃO DE DESEMPENHO ─────────────────────────────────
    banner(ws, 10, 'AVALIAÇÃO DE DESEMPENHO')
    for i, h in enumerate(['Critério de Avaliação', 'PESO (a)', 'NOTA (b)\n(0, 1 ou 3)',
                           'SUBTOTAL\n(c = a × b)', 'Status']):
        put(ws, '%s11' % get_column_letter(1 + i), h, font=F_HDR, fill=FILL_HDR,
            align='center', wrap=True)
    merge(ws, 'F11', 'J11', fill=FILL_HDR)
    put(ws, 'F11', 'Aviso de Justificativa', font=F_HDR, fill=FILL_HDR, align='center')
    ws.row_dimensions[11].height = 28

    dv = DataValidation(type='list', formula1='"0,1,3"', allow_blank=False,
                        showErrorMessage=True, errorTitle='Nota inválida',
                        error='A nota deve ser 0, 1 ou 3.')
    ws.add_data_validation(dv)

    r, grupo_tot = 12, {}
    for grp in (1, 2, 3):
        banner(ws, r, GRUPO_NOME[grp], fill=FILL_GRP, font=F_GRP)
        r += 1
        ini = r
        for idx, (cid, nome, peso, g, adm) in enumerate(CRITERIOS):
            if g != grp:
                continue
            fr = FILL_ADM if adm else None
            put(ws, 'A%d' % r, '=Dados!$B$%d' % (pz0 + idx), font=F_TXT, fill=fr, wrap=True)
            put(ws, 'B%d' % r, '=Dados!$C$%d' % (pz0 + idx), font=F_TXT, fill=fr,
                fmt=NUM2, align='center')
            put(ws, 'C%d' % r, 3, font=F_IN, fill=FILL_IN, align='center', unlocked=True)
            dv.add(ws['C%d' % r])
            put(ws, 'D%d' % r, '=B{r}*C{r}'.format(r=r), font=F_B, fill=FILL_CAL,
                fmt=NUM2, align='center')
            put(ws, 'E%d' % r, '=IF(OR(C{r}=0,C{r}=1,C{r}=3),"✓ Ok","⛔ Inválida")'.format(r=r),
                font=F_TXT, fill=fr, align='center')
            merge(ws, 'F%d' % r, 'J%d' % r, fill=fr)
            put(ws, 'F%d' % r,
                '=IF(OR(C{r}=0,C{r}=1),"⚠ Favor inserir justificativa da nota "&C{r}&'
                '" no campo JUSTIFICATIVAS","")'.format(r=r), font=F_RED, fill=fr, wrap=True)
            r += 1
        fim = r - 1
        put(ws, 'A%d' % r, 'TOTAL GRUPO %d' % grp, font=F_B, fill=FILL_TOT)
        put(ws, 'B%d' % r, '=SUM(B%d:B%d)' % (ini, fim), font=F_B, fill=FILL_TOT,
            fmt=NUM2, align='center')
        put(ws, 'C%d' % r, None, fill=FILL_TOT)
        put(ws, 'D%d' % r, '=SUM(D%d:D%d)' % (ini, fim), font=F_B, fill=FILL_CAL,
            fmt=NUM2, align='center')
        put(ws, 'E%d' % r, None, fill=FILL_TOT)
        merge(ws, 'F%d' % r, 'J%d' % r, fill=FILL_TOT)
        grupo_tot[grp] = 'D%d' % r
        r += 1

    # ── NOTA FINAL / CONCEITO ───────────────────────────────────
    rnf = r
    nota_ini, nota_fim = 13, rnf - 1          # faixa que contém as notas (coluna C)
    merge(ws, 'A%d' % rnf, 'C%d' % rnf, fill=FILL_TIT)
    put(ws, 'A%d' % rnf, 'NOTA FINAL (somatório dos Grupos 1, 2 e 3)', font=F_TIT,
        fill=FILL_TIT)
    NF = 'D%d' % rnf
    put(ws, NF, '=%s+%s+%s' % (grupo_tot[1], grupo_tot[2], grupo_tot[3]),
        font=Font(name=FONTE, size=13, bold=True, color='FFFFFF'), fill=FILL_TIT,
        fmt=NUM2, align='center')
    put(ws, 'E%d' % rnf, 'Conceito:', font=F_TIT, fill=FILL_TIT, align='right')
    CONC = 'F%d' % rnf
    put(ws, CONC, '=INDEX(Dados!$B$%d:$B$%d,MATCH(%s,Dados!$A$%d:$A$%d,1))'
        % (fx0, fx1, NF, fx0, fx1),
        font=Font(name=FONTE, size=13, bold=True, color='FFFFFF'), fill=FILL_TIT,
        align='center')
    merge(ws, 'G%d' % rnf, 'J%d' % rnf, fill=FILL_TIT)
    put(ws, 'G%d' % rnf, '(¹) Grupo 3 – Salários/Benefícios/Obrigações Trabalhistas: '
                         'preenchido pelo Fiscal Administrativo',
        font=Font(name=FONTE, size=8, color='FFFFFF'), fill=FILL_TIT, wrap=True)
    ws.row_dimensions[rnf].height = 24

    # ── POSTOS DE TRABALHO ──────────────────────────────────────
    rpb = rnf + 2
    banner(ws, rpb, 'POSTOS DE TRABALHO')
    rph = rpb + 1
    for i, h in enumerate(['Item', 'POSTO', 'Qtde\nContratada',
                           'Valor Unit.\nContratado (R$)', 'Horas\nContratadas',
                           'Sem\nCobertura (h)', 'Valor da Hora\nContratada (R$)',
                           'Valor do Item após\nDescontos (R$)',
                           'Mensal\nContratado (R$)', '']):
        put(ws, '%s%d' % (get_column_letter(1 + i), rph), h, font=F_HDR, fill=FILL_HDR,
            align='center', wrap=True)
    ws.row_dimensions[rph].height = 32
    p0 = rph + 1
    p1 = p0 + NP - 1
    rpt = p0 + NP                              # linha de totais dos postos
    SOMA_DESC, SOMA_MENS = 'H%d' % rpt, 'I%d' % rpt

    # ── posições do RESUMO FINANCEIRO (o IC fica entre os dois blocos) ──
    ric = rpt + 2                              # banner do instrumento de cobrança
    rres = ric + 2 + NP + 1                    # banner do resumo financeiro
    PCTF     = '$F$%d' % (rres + 4)            # % do faturamento a ser liberado
    GLOSA    = 'F%d' % (rres + 5)
    DIAS_MES = '$F$%d' % (rres + 6)
    DIAS_IN  = 'F%d' % (rres + 7)
    DIAS_FIN = '$F$%d' % (rres + 8)

    for i in range(NP):
        rr = p0 + i
        dr = D['posto_row0'] + i
        put(ws, 'A%d' % rr, '=IF(Dados!$A$%d="","",Dados!$A$%d)' % (dr, dr),
            font=F_LINK, align='center')
        put(ws, 'B%d' % rr, '=IF(Dados!$B$%d="","",Dados!$B$%d)' % (dr, dr), font=F_LINK)
        put(ws, 'C%d' % rr, '=IF(Dados!$C$%d="","",Dados!$C$%d)' % (dr, dr),
            font=F_LINK, align='center')
        put(ws, 'D%d' % rr, '=IF(Dados!$D$%d="","",Dados!$D$%d)' % (dr, dr),
            font=F_LINK, fmt=MOEDA, align='center')
        put(ws, 'E%d' % rr, '=IF(C{r}="","",C{r}*{h})'.format(r=rr, h=HORAS_MES),
            font=F_TXT, fill=FILL_CAL, fmt=INT, align='center')
        put(ws, 'F%d' % rr, None, font=F_IN, fill=FILL_IN, align='center', unlocked=True)
        put(ws, 'G%d' % rr, '=IF(C{r}="","",IF(C{r}>0,ROUNDDOWN(N(D{r})/{h},2),0))'.format(
            r=rr, h=HORAS_MES), font=F_TXT, fill=FILL_CAL, fmt=MOEDA, align='center')
        put(ws, 'H%d' % rr,
            '=IF(C{r}="","",IF(AND(C{r}>0,N(D{r})>0),MAX(0,ROUNDDOWN((C{r}*N(D{r})'
            '-G{r}*N(F{r}))/30*{d}*{p},2)),0))'.format(r=rr, d=DIAS_FIN, p=PCTF),
            font=F_B, fill=FILL_CAL, fmt=MOEDA, align='center')
        put(ws, 'I%d' % rr, '=IF(C{r}="","",ROUNDDOWN(C{r}*N(D{r}),2))'.format(r=rr),
            font=F_TXT, fill=FILL_CAL, fmt=MOEDA, align='center')
        put(ws, 'J%d' % rr, None)

    put(ws, 'A%d' % rpt, 'TOTAIS', font=F_B, fill=FILL_TOT)
    for L in 'BCDEGJ':
        put(ws, '%s%d' % (L, rpt), None, fill=FILL_TOT)
    put(ws, 'F%d' % rpt, '=SUM(F%d:F%d)' % (p0, p1), font=F_B, fill=FILL_TOT,
        fmt=INT, align='center')
    put(ws, SOMA_DESC, '=SUM(H%d:H%d)' % (p0, p1), font=F_B, fill=FILL_TOT,
        fmt=MOEDA, align='center')
    put(ws, SOMA_MENS, '=SUM(I%d:I%d)' % (p0, p1), font=F_B, fill=FILL_TOT,
        fmt=MOEDA, align='center')

    # ── INSTRUMENTO DE COBRANÇA ─────────────────────────────────
    banner(ws, ric, '📋 INSTRUMENTO DE COBRANÇA',
           fill=PatternFill('solid', fgColor=VERDE_ESC))
    put(ws, 'A%d' % (ric + 1), 'Item', font=F_HDR, fill=FILL_ICH, align='center')
    merge(ws, 'B%d' % (ric + 1), 'E%d' % (ric + 1), fill=FILL_ICH)
    put(ws, 'B%d' % (ric + 1), 'Posto', font=F_HDR, fill=FILL_ICH, align='center')
    merge(ws, 'F%d' % (ric + 1), 'G%d' % (ric + 1), fill=FILL_ICH)
    put(ws, 'F%d' % (ric + 1), 'Qtde Faturada', font=F_HDR, fill=FILL_ICH, align='center')
    merge(ws, 'H%d' % (ric + 1), 'J%d' % (ric + 1), fill=FILL_ICH)
    put(ws, 'H%d' % (ric + 1), 'Valor Unit. Faturado (R$)', font=F_HDR, fill=FILL_ICH,
        align='center')
    for i in range(NP):
        rr = ric + 2 + i
        pr = p0 + i
        put(ws, 'A%d' % rr, '=IF(A%d="","",A%d)' % (pr, pr), font=F_TXT, fill=FILL_IC,
            align='center')
        merge(ws, 'B%d' % rr, 'E%d' % rr, fill=FILL_IC)
        put(ws, 'B%d' % rr, '=IF(B%d="","",B%d)' % (pr, pr), font=F_TXT, fill=FILL_IC)
        merge(ws, 'F%d' % rr, 'G%d' % rr, fill=FILL_IC)
        put(ws, 'F%d' % rr, '=IF(C%d="","",C%d)' % (pr, pr), font=F_TXT, fill=FILL_IC,
            fmt=INT, align='center')
        merge(ws, 'H%d' % rr, 'J%d' % rr, fill=FILL_IC)
        put(ws, 'H%d' % rr, '=IF(C{p}="","",IF(C{p}>0,ROUNDDOWN(H{p}/C{p},2),0))'.format(p=pr),
            font=F_B, fill=FILL_IC, fmt=MOEDA, align='center')

    # ── RESUMO FINANCEIRO ───────────────────────────────────────
    banner(ws, rres, 'RESUMO FINANCEIRO')
    linhas = [
        ('Contratado Mensal – Total (R$)   (Σ Qtde × Valor Unitário de todos os postos)',
         '=%s' % SOMA_MENS, MOEDA, FILL_CAL, F_B, ''),
        ('Nota Final Apurada', '=%s' % NF, NUM2, FILL_CAL, F_B, ''),
        ('Nota Atribuída pelo Avaliador (Conceito)', '=%s' % CONC, None, FILL_CAL,
         Font(name=FONTE, size=12, bold=True, color=AZUL_ESC), ''),
        ('% do Faturamento a ser Liberado',
         '=INDEX(Dados!$C$%d:$C$%d,MATCH(%s,Dados!$A$%d:$A$%d,1))' % (fx0, fx1, NF, fx0, fx1),
         PCT, FILL_CAL, F_B, ''),
        ('Valor Glosado por QUALIDADE (R$)',
         '=ROUNDDOWN({m}-{m}*{p},2)'.format(m=SOMA_MENS, p=PCTF), MOEDA, FILL_RED,
         Font(name=FONTE, size=10, bold=True, color='B71C1C'), ''),
        ('Número de Dias do Mês (calculado)',
         '=IF(ISNUMBER(B7),DAY(EOMONTH(B7,0)),30)', INT, FILL_CAL, F_B, ''),
        ('Número de Dias Trabalhados no Mês',
         '=IF(ISNUMBER(B7),DAY(EOMONTH(B7,0)),30)', INT, FILL_IN, F_IN,
         'Campo de preenchimento: altere apenas em mês parcial (início/fim de contrato)'),
        ('Dias para Cálculo (mês cheio = 30 avos)',
         '=IF(N({d})>={dm},30,MAX(0,N({d})))'.format(d=DIAS_IN, dm=DIAS_MES),
         NUM2, FILL_CAL, F_B, ''),
        ('Soma dos Itens após Descontos (R$)', '=%s' % SOMA_DESC, MOEDA, FILL_CAL, F_B, ''),
    ]
    for i, (lbl, formula, fmt, fill, font, obs) in enumerate(linhas):
        rr = rres + 1 + i
        rotulo(ws, rr, lbl)
        merge(ws, 'F%d' % rr, 'G%d' % rr, fill=fill)
        put(ws, 'F%d' % rr, formula, font=font, fill=fill, fmt=fmt, align='center',
            unlocked=(fill is FILL_IN))
        merge(ws, 'H%d' % rr, 'J%d' % rr)
        put(ws, 'H%d' % rr, obs or None, font=F_SM, wrap=True)

    # ── REEMBOLSO DE CRECHE (definido antes do Valor Aprovado) ──
    rap = rres + len(linhas) + 1               # valor aprovado
    rcr = rap + 2                              # banner creche
    CR_RET, CR_APR = 'F%d' % (rcr + 1), 'F%d' % (rcr + 2)
    CR_UNI, CR_TOT, CR_SAL = 'F%d' % (rcr + 3), 'F%d' % (rcr + 4), 'F%d' % (rcr + 5)

    rotulo(ws, rap, 'VALOR APROVADO PARA PAGAMENTO (R$)   (Σ Itens após Descontos + '
                    'Total dos Reembolsos de Creche)')
    ws['A%d' % rap].font = F_TIT
    for row in ws['A%d:E%d' % (rap, rap)]:
        for c in row:
            c.fill = FILL_TIT
    merge(ws, 'F%d' % rap, 'G%d' % rap, fill=FILL_APR)
    APROVADO = 'F%d' % rap
    put(ws, APROVADO, '=MAX(0,ROUNDDOWN(%s+%s,2))' % (SOMA_DESC, CR_TOT),
        font=Font(name=FONTE, size=12, bold=True, color='1B5E20'), fill=FILL_APR,
        fmt=MOEDA, align='center')
    merge(ws, 'H%d' % rap, 'J%d' % rap)
    put(ws, 'H%d' % rap, None)
    ws.row_dimensions[rap].height = 26

    banner(ws, rcr, '🏫 REEMBOLSO DE CRECHE', fill=PatternFill('solid', fgColor='F57F17'))
    prev = ("'%s'!%s" % (aba_anterior, CR_SAL)) if aba_anterior else None
    creche = [
        ('Valor Reembolso Creche Retido (R$)', '=Dados!$%s' % D['creche_retido'],
         MOEDA, FILL_IN, F_IN),
        ('Reembolsos Aprovados (quantidade)', None, INT, FILL_IN, F_IN),
        ('Reembolso por dependente (R$)', '=Dados!$%s' % D['creche_unit'],
         MOEDA, FILL_IN, F_IN),
        ('Valor Total dos Reembolsos (R$)   (Aprovados × Reembolso por dependente)',
         '=ROUNDDOWN(N(%s)*N(%s),2)' % (CR_APR, CR_UNI), MOEDA, FILL_CAL, F_B),
        ('Saldo de Reembolso (R$)   (acumulado: saldo da aba anterior + Retido – Total)',
         '=ROUNDDOWN(%sN(%s)-N(%s),2)' % (('%s+' % prev) if prev else '', CR_RET, CR_TOT),
         MOEDA, FILL_CAL, F_B),
    ]
    for i, (lbl, formula, fmt, fill, font) in enumerate(creche):
        rr = rcr + 1 + i
        rotulo(ws, rr, lbl)
        merge(ws, 'F%d' % rr, 'G%d' % rr, fill=fill)
        put(ws, 'F%d' % rr, formula, font=font, fill=fill, fmt=fmt, align='center',
            unlocked=(fill is FILL_IN))
        merge(ws, 'H%d' % rr, 'J%d' % rr)
        put(ws, 'H%d' % rr, 'Campo de preenchimento' if fill is FILL_IN else None,
            font=F_SM, wrap=True)
    rwarn = rcr + 6
    merge(ws, 'A%d' % rwarn, 'J%d' % rwarn, fill=FILL_RED)
    put(ws, 'A%d' % rwarn, '=IF(%s<0,"⚠ ATENÇÃO: Saldo negativo — será necessário ADITIVAR '
                           'O CONTRATO PARA SUPORTAR A DESPESA!","")' % CR_SAL,
        font=F_RED, fill=FILL_RED, wrap=True)

    # ── FAIXA DE AJUSTE DE PAGAMENTO ────────────────────────────
    rfx = rwarn + 2
    banner(ws, rfx, 'FAIXA DE AJUSTE DE PAGAMENTO')
    merge(ws, 'A%d' % (rfx + 1), 'C%d' % (rfx + 1), fill=FILL_HDR)
    put(ws, 'A%d' % (rfx + 1), 'Conceito', font=F_HDR, fill=FILL_HDR, align='center')
    merge(ws, 'D%d' % (rfx + 1), 'F%d' % (rfx + 1), fill=FILL_HDR)
    put(ws, 'D%d' % (rfx + 1), '% do Faturamento a ser Pago', font=F_HDR, fill=FILL_HDR,
        align='center')
    merge(ws, 'G%d' % (rfx + 1), 'J%d' % (rfx + 1), fill=FILL_HDR)
    put(ws, 'G%d' % (rfx + 1), 'Nota Média Mensal Apurada', font=F_HDR, fill=FILL_HDR,
        align='center')
    for i in range(len(FAIXA)):
        rr = rfx + 2 + i
        dr = fx1 - i                            # exibe de A (maior nota) até F
        merge(ws, 'A%d' % rr, 'C%d' % rr)
        put(ws, 'A%d' % rr, '=Dados!$B$%d' % dr, font=F_B, align='center')
        merge(ws, 'D%d' % rr, 'F%d' % rr)
        put(ws, 'D%d' % rr, '=Dados!$C$%d' % dr, font=F_TXT, fmt=PCT, align='center')
        merge(ws, 'G%d' % rr, 'J%d' % rr)
        put(ws, 'G%d' % rr, '=Dados!$D$%d' % dr, font=F_TXT, align='center')
    ws.conditional_formatting.add(
        'A%d:J%d' % (rfx + 2, rfx + 1 + len(FAIXA)),
        FormulaRule(formula=['$A%d=$F$%d' % (rfx + 2, rnf)],
                    fill=PatternFill('solid', fgColor='FFF59D'),
                    font=Font(name=FONTE, bold=True)))

    # ── ASSINATURAS ─────────────────────────────────────────────
    rass = rfx + 2 + len(FAIXA) + 1
    banner(ws, rass, 'APROVAÇÃO E ASSINATURAS')
    blocos = [('Fiscal Operacional', 'A', 'C', 'Nome completo', 'Matrícula'),
              ('Fiscal Administrativo (¹)', 'D', 'F', 'Nome completo', 'Matrícula'),
              ('Representante da Contratada', 'G', 'H', 'Nome completo', 'Cargo'),
              ('Data da Aprovação', 'I', 'J', 'dd/mm/aaaa', '')]
    for lbl, c1, c2, ph1, ph2 in blocos:
        merge(ws, '%s%d' % (c1, rass + 1), '%s%d' % (c2, rass + 1), fill=FILL_TOT)
        put(ws, '%s%d' % (c1, rass + 1), lbl, font=F_B, fill=FILL_TOT, align='center',
            wrap=True)
        merge(ws, '%s%d' % (c1, rass + 2), '%s%d' % (c2, rass + 2), fill=FILL_IN)
        put(ws, '%s%d' % (c1, rass + 2), None, font=F_IN, fill=FILL_IN, align='center',
            unlocked=True, fmt=(DATA_BR if c1 == 'I' else None))
        merge(ws, '%s%d' % (c1, rass + 3), '%s%d' % (c2, rass + 3),
              fill=(FILL_IN if ph2 else None))
        put(ws, '%s%d' % (c1, rass + 3), None, font=F_IN,
            fill=(FILL_IN if ph2 else None), align='center', unlocked=bool(ph2))
        merge(ws, '%s%d' % (c1, rass + 4), '%s%d' % (c2, rass + 4))
        put(ws, '%s%d' % (c1, rass + 4), '%s / %s' % (ph1, ph2) if ph2 else ph1,
            font=F_SM, align='center', wrap=True)
    ws.row_dimensions[rass + 2].height = 20
    ws.row_dimensions[rass + 3].height = 18

    # ── JUSTIFICATIVAS ──────────────────────────────────────────
    rjus = rass + 6
    banner(ws, rjus, '📝 JUSTIFICATIVAS (obrigatório para notas 0 e 1)')
    merge(ws, 'A%d' % (rjus + 1), 'J%d' % (rjus + 4), fill=FILL_IN)
    c = put(ws, 'A%d' % (rjus + 1), None, font=F_IN, fill=FILL_IN, unlocked=True)
    c.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
    for k in range(1, 5):
        ws.row_dimensions[rjus + k].height = 22
    merge(ws, 'A%d' % (rjus + 5), 'J%d' % (rjus + 5))
    put(ws, 'A%d' % (rjus + 5),
        '=IF(COUNTIF(C{a}:C{b},0)+COUNTIF(C{a}:C{b},1)>0,"⚠ Existem notas 0 ou 1 nesta aba '
        '— a justificativa acima é OBRIGATÓRIA. Informe critério, nota e fundamentação.",'
        '"Para cada nota 0 ou 1 informe: critério, nota atribuída e fundamentação.")'.format(
            a=nota_ini, b=nota_fim), font=F_RED, wrap=True)

    # ── IMPRESSÃO: 1 folha A4 (equivale ao fitToPage do formulário) ──
    ws.print_area = 'A1:J%d' % (rjus + 5)
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.page_margins.top = ws.page_margins.bottom = 0.3
    ws.freeze_panes = 'A12'

    return {'periodo': 'B7', 'nf': NF, 'conc': CONC, 'pct': PCTF.replace('$', ''),
            'mens': SOMA_MENS, 'glosa': GLOSA, 'desc': SOMA_DESC, 'aprov': APROVADO,
            'creche_tot': CR_TOT, 'creche_sal': CR_SAL, 'dias': DIAS_IN}


# ═════════════════════════════════════════════════════════════════
# ABA RESUMO GERAL – consolida os 60 periodos (equivale ao CSV)
# ═════════════════════════════════════════════════════════════════
def aba_resumo(wb, refs, nomes):
    ws = wb.create_sheet('Resumo Geral', 1)
    ws.sheet_view.showGridLines = False
    cols = [('Aba', 8), ('Período', 16), ('Nota Final', 11), ('Conceito', 10),
            ('% Liberado', 11), ('Dias Trab.', 10), ('Contratado Mensal (R$)', 19),
            ('Glosa por Qualidade (R$)', 19), ('Itens após Descontos (R$)', 20),
            ('Reembolsos Creche (R$)', 19), ('Saldo Creche (R$)', 17),
            ('Valor Aprovado (R$)', 19)]
    banner(ws, 1, 'RESUMO GERAL DOS 60 PERÍODOS', last=get_column_letter(len(cols)))
    for i, (h, w) in enumerate(cols):
        L = get_column_letter(1 + i)
        ws.column_dimensions[L].width = w
        put(ws, '%s2' % L, h, font=F_HDR, fill=FILL_HDR, align='center', wrap=True)
    ws.row_dimensions[2].height = 32
    for i, (nome, rf) in enumerate(zip(nomes, refs)):
        r, q = 3 + i, "'%s'!" % nome
        put(ws, 'A%d' % r, nome, font=F_B, align='center')
        put(ws, 'B%d' % r, '=%s%s' % (q, rf['periodo']), font=F_TXT, fmt=MES_ANO,
            align='center')
        put(ws, 'C%d' % r, '=%s%s' % (q, rf['nf']), font=F_TXT, fmt=NUM2, align='center')
        put(ws, 'D%d' % r, '=%s%s' % (q, rf['conc']), font=F_B, align='center')
        put(ws, 'E%d' % r, '=%s%s' % (q, rf['pct']), font=F_TXT, fmt=PCT, align='center')
        put(ws, 'F%d' % r, '=%s%s' % (q, rf['dias']), font=F_TXT, fmt=INT, align='center')
        put(ws, 'G%d' % r, '=%s%s' % (q, rf['mens']), font=F_TXT, fmt=MOEDA)
        put(ws, 'H%d' % r, '=%s%s' % (q, rf['glosa']), font=F_TXT, fmt=MOEDA)
        put(ws, 'I%d' % r, '=%s%s' % (q, rf['desc']), font=F_TXT, fmt=MOEDA)
        put(ws, 'J%d' % r, '=%s%s' % (q, rf['creche_tot']), font=F_TXT, fmt=MOEDA)
        put(ws, 'K%d' % r, '=%s%s' % (q, rf['creche_sal']), font=F_TXT, fmt=MOEDA)
        put(ws, 'L%d' % r, '=%s%s' % (q, rf['aprov']), font=F_B, fill=FILL_CAL, fmt=MOEDA)
    rt = 3 + NUM_PERIODS
    merge(ws, 'A%d' % rt, 'F%d' % rt, fill=FILL_TOT)
    put(ws, 'A%d' % rt, 'TOTAIS DO CONTRATO', font=F_B, fill=FILL_TOT)
    for L in 'GHIJL':
        put(ws, '%s%d' % (L, rt), '=SUM(%s3:%s%d)' % (L, L, rt - 1), font=F_B,
            fill=FILL_TOT, fmt=MOEDA)
    put(ws, 'K%d' % rt, '=K%d' % (rt - 1), font=F_B, fill=FILL_TOT, fmt=MOEDA)
    merge(ws, 'A%d' % (rt + 1), 'L%d' % (rt + 1))
    put(ws, 'A%d' % (rt + 1), 'Saldo de Creche: valor acumulado do último período (não é '
                              'soma). Linhas em vermelho indicam saldo negativo — exige '
                              'aditivo contratual.', font=F_SM, wrap=True)
    ws.conditional_formatting.add(
        'K3:K%d' % (rt - 1),
        FormulaRule(formula=['K3<0'], fill=FILL_RED,
                    font=Font(name=FONTE, bold=True, color='B71C1C')))
    ws.freeze_panes = 'C3'


# ═════════════════════════════════════════════════════════════════
# ABA INSTRUCOES
# ═════════════════════════════════════════════════════════════════
def aba_instrucoes(wb, contratada, contrato, unidade):
    ws = wb.create_sheet('Instruções', 0)
    ws.column_dimensions['A'].width = 3
    ws.column_dimensions['B'].width = 44
    ws.column_dimensions['C'].width = 82
    ws.sheet_view.showGridLines = False
    put(ws, 'B1', 'IMR %s – Contrato %s – %s' % (contratada, contrato, unidade),
        font=Font(name=FONTE, size=14, bold=True, color=AZUL_ESC), border=False)
    put(ws, 'B2', 'Instrumento de Medição de Resultado – versão Excel, 60 períodos '
                  '(equivalente ao formulário web da PR-6)',
        font=Font(name=FONTE, size=10, italic=True, color='555555'), border=False)

    secoes = [
        ('ESTRUTURA DA PASTA DE TRABALHO', [
            ('Dados', 'Contrato, unidade, contratada, fiscal, período inicial, postos '
                      'contratados, valores de creche, pesos dos critérios e faixa de '
                      'ajuste. Tudo o que é preenchido aqui vale para as 60 abas.'),
            ('Resumo Geral', 'Consolida período, nota, conceito, %, glosa, creche e valor '
                             'aprovado dos 60 períodos (equivale à exportação CSV do '
                             'formulário web).'),
            ('P01 … P60', 'Uma aba por período (mês), com avaliação de desempenho, postos '
                          'de trabalho, instrumento de cobrança, resumo financeiro, '
                          'reembolso de creche, assinaturas e justificativas.'),
        ]),
        ('COMO PREENCHER', [
            ('1. Aba Dados – identificação',
             'Contrato, unidade, contratada, fiscal operacional e o PERÍODO INICIAL '
             '(ex.: 01/05/2024). As abas P01…P60 recebem os meses seguintes '
             'automaticamente, como a propagação de períodos do formulário web.'),
            ('2. Aba Dados – postos',
             'Liste os postos com Item, Nome, Qtde Contratada e Valor Unitário Contratado. '
             'Eles são replicados nas 60 abas. A linha 12 traz um exemplo: substitua-a.'),
            ('3. Em cada aba Pxx',
             'Escolha a NOTA de cada critério (lista suspensa 0, 1 ou 3), informe as HORAS '
             'SEM COBERTURA de cada posto, os DIAS TRABALHADOS (apenas em mês parcial) e '
             'os dados de creche do mês.'),
            ('4. Justificativas',
             'Notas 0 ou 1 geram aviso na linha do critério e ao final da aba; a '
             'justificativa é obrigatória.'),
            ('5. Assinaturas', 'Nome e matrícula dos fiscais, nome e cargo do '
                               'representante da contratada e data da aprovação.'),
        ]),
        ('CÁLCULOS – idênticos ao formulário web', [
            ('Subtotal do critério', 'Peso (a) × Nota (b)'),
            ('Nota Final', 'Soma dos totais dos Grupos 1, 2 e 3 (máximo 9,00)'),
            ('Conceito e % liberado',
             'Faixa de ajuste: A = 100% (9,00 a 8,10), B = 95% (8,09 a 7,65), '
             'C = 90% (7,64 a 6,75), D = 85% (6,74 a 5,85), E = 80% (5,84 a 5,40), '
             'F = 75% (abaixo de 5,40)'),
            ('Horas Contratadas', 'Qtde Contratada × 220'),
            ('Valor da Hora Contratada', 'Valor Unitário ÷ 220, truncado em 2 casas'),
            ('Valor do Item após Descontos',
             '((Qtde × Valor Unit.) − Valor da Hora × Horas sem Cobertura) ÷ 30 × Dias '
             'para Cálculo × % do Faturamento — truncado em 2 casas e nunca negativo'),
            ('Dias para Cálculo',
             'Mês cheio = 30 avos (inclusive em fevereiro e meses de 31 dias); mês '
             'parcial = dias trabalhados informados'),
            ('Mensal Contratado (por posto)',
             'Qtde × Valor Unitário, truncado em 2 casas (coluna de conferência; a soma é '
             'o Contratado Mensal do resumo financeiro)'),
            ('Valor Unit. Faturado (instrumento de cobrança)',
             'Valor do Item após Descontos ÷ Qtde Faturada, truncado em 2 casas'),
            ('Valor Glosado por Qualidade',
             'Contratado Mensal − (Contratado Mensal × % do Faturamento)'),
            ('Valor Aprovado para Pagamento',
             'Σ Valor dos Itens após Descontos + Valor Total dos Reembolsos de Creche'),
            ('Total dos Reembolsos de Creche',
             'Reembolsos Aprovados × Reembolso por dependente, truncado em 2 casas'),
            ('Saldo de Reembolso de Creche',
             'Acumulado: saldo da aba anterior + Retido − Total dos Reembolsos. Saldo '
             'negativo exibe o aviso de aditivo contratual.'),
        ]),
        ('LEGENDA DE CORES', [
            ('Fundo amarelo com texto azul', 'Campo de preenchimento manual'),
            ('Fundo verde claro', 'Valor calculado por fórmula – não editar'),
            ('Texto verde', 'Valor trazido da aba Dados'),
            ('Azul escuro / azul', 'Títulos e cabeçalhos'),
            ('Vermelho', 'Glosa, saldo negativo e avisos de justificativa'),
            ('Bege (Grupo 3)', 'Critério preenchido pelo Fiscal Administrativo (¹)'),
        ]),
        ('EQUIVALÊNCIA COM OS RECURSOS DO FORMULÁRIO WEB', [
            ('Travar formulário (senha)',
             'Revisão ▸ Proteger Planilha, com senha. Os campos de preenchimento já estão '
             'desbloqueados: a proteção trava apenas fórmulas, postos e dados do contrato.'),
            ('Salvamento automático e sincronização entre computadores',
             'Salve a pasta no OneDrive, SharePoint ou Google Drive: o salvamento '
             'automático e a edição simultânea passam a valer para todos os usuários.'),
            ('Salvar PDF da aba', 'Arquivo ▸ Exportar ▸ PDF. Cada aba já está configurada '
                                  'para caber em 1 folha A4.'),
            ('Exportação CSV', 'Aba Resumo Geral ▸ Arquivo ▸ Salvar como ▸ CSV.'),
            ('Validação das notas', 'Lista suspensa com 0, 1 e 3 em cada critério; outros '
                                    'valores são recusados.'),
        ]),
    ]
    r = 4
    for titulo, itens in secoes:
        merge(ws, 'B%d' % r, 'C%d' % r, fill=FILL_TIT)
        put(ws, 'B%d' % r, titulo, font=F_TIT, fill=FILL_TIT)
        ws.row_dimensions[r].height = 18
        r += 1
        for a, b in itens:
            put(ws, 'B%d' % r, a, font=F_B, fill=FILL_TOT, wrap=True, valign='top')
            put(ws, 'C%d' % r, b, font=F_TXT, wrap=True, valign='top')
            ws.row_dimensions[r].height = 14 + 11 * (len(b) // 100)
            r += 1
        r += 1
    merge(ws, 'B%d' % r, 'C%d' % r)
    put(ws, 'B%d' % r, 'Gerado a partir do formulário "Modelo IMR Vigilância Definitivo V2 '
                       'Com CRECHE – PR6" (versão 60 períodos) por gen_imr_excel.py.',
        font=F_SM, wrap=True)


# ═════════════════════════════════════════════════════════════════
def ajusta_impressao(ws, ultima_coluna, ultima_linha):
    ws.print_area = 'A1:%s%d' % (ultima_coluna, ultima_linha)
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = '1:2'


def main():
    contratada = sys.argv[1] if len(sys.argv) > 1 else 'Angels'
    contrato   = sys.argv[2] if len(sys.argv) > 2 else '04/24'
    unidade    = sys.argv[3] if len(sys.argv) > 3 else 'CBAE'
    saida      = sys.argv[4] if len(sys.argv) > 4 else 'IMR Angels - 04-24 - CBAE.xlsx'

    wb = Workbook()
    wb.remove(wb.active)
    D = aba_dados(wb, contratada, contrato, unidade)

    refs, nomes, anterior = [], [], None
    for n in range(1, NUM_PERIODS + 1):
        nome = 'P%02d' % n
        refs.append(aba_periodo(wb, n, D, nome, anterior))
        nomes.append(nome)
        anterior = nome
    aba_resumo(wb, refs, nomes)        # entra no índice 1
    aba_instrucoes(wb, contratada, contrato, unidade)   # entra no índice 0
    ajusta_impressao(wb['Instruções'], 'C', wb['Instruções'].max_row)
    ajusta_impressao(wb['Dados'], 'F', wb['Dados'].max_row)
    ajusta_impressao(wb['Resumo Geral'], 'L', wb['Resumo Geral'].max_row)
    wb.active = 0
    wb.save(saida)
    print('OK ->', saida)


if __name__ == '__main__':
    main()
